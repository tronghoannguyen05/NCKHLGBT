"""Dựng bản Word của bài tiếng Việt theo thể lệ của Tạp chí Nghiên cứu Con người, Gia đình và Giới.

Nguồn   docs/bai_tieng_viet/bai_bao.md
Đầu ra  docs/bai_tieng_viet/Bai_bao.docx

Thể lệ: khổ A4, Times New Roman 12, lề trên 2,5 cm, lề dưới 2,5 cm, lề trái 3 cm, lề phải 2 cm; tên bảng
đặt bên trên, chữ đậm, căn giữa; nguồn đặt bên dưới, chữ nghiêng, căn giữa; bảng in đen trắng.

Chạy: python3 docs/bai_tieng_viet/build_docx.py   (cần pandoc hoặc pypandoc_binary, và python-docx)
Sửa nội dung ở bai_bao.md rồi chạy lại, không sửa trực tiếp tệp Word.
"""
import os
import re
import shutil
import subprocess
import tempfile

from docx import Document
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_LINE_SPACING
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Pt

HERE = os.path.dirname(os.path.abspath(__file__))
SRC = os.path.join(HERE, 'bai_bao.md')
OUT = os.path.join(HERE, 'Bai_bao.docx')
TEXT_WIDTH = 9072  # twips: 21 cm - 3 cm - 2 cm = 16 cm

# Tỷ lệ độ rộng cột và số cột căn trái của từng bảng
TABLES = {
    '1': ([0.80, 0.20], 1),
    '2': ([0.75, 0.25], 1),
    '3': ([0.52, 0.12, 0.18, 0.18], 1),
    '4': ([0.34, 0.22, 0.22, 0.22], 1),
    '5': ([0.52, 0.16, 0.16, 0.16], 1),
    '6': ([0.38, 0.10, 0.13, 0.13, 0.13, 0.13], 1),
    '7': ([0.46, 0.18, 0.18, 0.18], 1),
    '8': ([0.28, 0.08, 0.17, 0.15, 0.16, 0.16], 1),
    '9': ([0.14, 0.38, 0.10, 0.18, 0.10, 0.10], 2),
    '10': ([0.24, 0.36, 0.40], 3),
    'PL1': ([0.36, 0.07, 0.19, 0.19, 0.19], 1),
}


def styled(style, text):
    return '::: {custom-style="' + style + '"}\n' + text.strip() + '\n:::'


def pandoc_path():
    try:
        import pypandoc
        return pypandoc.get_pandoc_path()
    except Exception:
        path = shutil.which('pandoc')
        if not path:
            raise SystemExit('pandoc not found')
        return path


def to_markdown(text):
    """Gắn kiểu cho tên bài, khối tác giả, tên bảng, ghi chú, nguồn và tài liệu tham khảo."""
    paras = [p.strip() for p in re.split(r'\n\s*\n', text) if p.strip()]
    out, table_ids, in_refs, in_front = [], [], False, True
    for p in paras:
        if p.startswith('# '):
            out.append(styled('Title', p[2:]))
            continue
        if in_front and not p.startswith('## '):
            if p.startswith('*') and p.endswith('*') and not p.startswith('**'):
                out.append(styled('Subtitle', p))
            else:
                out.append(styled('Author Block', p))
            continue
        if p.startswith('## '):
            in_front = False
            in_refs = p.startswith('## Tài liệu tham khảo')
            out.append(p)
            continue
        if in_refs:
            out.append(styled('Reference', re.sub(r'^(\d+)\. ', r'\1\\. ', p)))
            continue
        m = re.match(r'\*\*Bảng (PL\d+|\d+)\. ', p)
        if m:
            table_ids.append(m.group(1))
            out.append(styled('Table Caption', p))
            continue
        if p.startswith('*Nguồn:'):
            out.append(styled('Table Source', p))
            continue
        if p.startswith('*Ghi chú:*'):
            out.append(styled('Table Note', p))
            continue
        if p.startswith('**Từ khóa:**') or p.startswith('**Keywords:**'):
            out.append(styled('Keywords', p))
            continue
        out.append(p)
    return '\n\n'.join(out) + '\n', table_ids


def el(tag, **attrs):
    e = OxmlElement('w:' + tag)
    for k, v in attrs.items():
        e.set(qn('w:' + k), str(v))
    return e


def set_style(d, name, size=12, bold=None, italic=None, align=None, first=0.0, hanging=None,
              before=0, after=6, keep=False, spacing=1.2):
    try:
        st = d.styles[name]
    except KeyError:
        return
    st.font.name = 'Times New Roman'
    st.font.size = Pt(size)
    st.font.color.rgb = None
    if bold is not None:
        st.font.bold = bold
    if italic is not None:
        st.font.italic = italic
    pf = st.paragraph_format
    pf.space_before = Pt(before)
    pf.space_after = Pt(after)
    pf.line_spacing_rule = WD_LINE_SPACING.MULTIPLE
    pf.line_spacing = spacing
    if align is not None:
        pf.alignment = align
    if hanging is not None:
        pf.left_indent = Cm(hanging)
        pf.first_line_indent = Cm(-hanging)
    else:
        pf.left_indent = Cm(0)
        pf.first_line_indent = Cm(first)
    pf.keep_with_next = keep


def format_styles(d):
    J, C, L = WD_ALIGN_PARAGRAPH.JUSTIFY, WD_ALIGN_PARAGRAPH.CENTER, WD_ALIGN_PARAGRAPH.LEFT
    set_style(d, 'Normal', align=J)
    for name in ('Body Text', 'First Paragraph'):
        set_style(d, name, align=J, first=1.0)
    set_style(d, 'Compact', align=L, after=0, spacing=1.0)
    set_style(d, 'Title', size=14, bold=True, align=C, after=6)
    set_style(d, 'Subtitle', size=12, italic=True, align=C, after=12)
    set_style(d, 'Author Block', align=C, after=4)
    set_style(d, 'Heading 1', size=12, bold=True, align=L, before=12, after=6, keep=True)
    set_style(d, 'Heading 2', size=12, bold=True, italic=True, align=L, before=10, after=6, keep=True)
    set_style(d, 'Heading 3', size=12, bold=True, italic=True, align=L, before=6, after=6, keep=True)
    set_style(d, 'Keywords', align=J, after=12)
    set_style(d, 'Table Caption', bold=True, italic=False, align=C, before=12, after=6, keep=True)
    set_style(d, 'Table Note', size=11, align=J, before=4, after=2, keep=True, spacing=1.0)
    set_style(d, 'Table Source', size=11, italic=True, align=C, before=2, after=12, spacing=1.0)
    set_style(d, 'Reference', align=J, hanging=1.0, after=4)
    for name in ('Heading 1', 'Heading 2', 'Heading 3', 'Title', 'Subtitle'):
        try:
            d.styles[name].font.color.rgb = None
        except KeyError:
            pass


def format_table(table, num):
    widths, left_cols = TABLES[num]
    assert len(table.columns) == len(widths), (num, len(table.columns), len(widths))
    tw = [int(round(w * TEXT_WIDTH)) for w in widths]
    tbl = table._tbl
    pr = tbl.tblPr
    for tag in ('tblW', 'jc', 'tblBorders', 'tblLayout', 'tblStyle', 'tblLook'):
        for old in pr.findall(qn('w:' + tag)):
            pr.remove(old)
    pr.append(el('tblW', w=sum(tw), type='dxa'))
    pr.append(el('jc', val='center'))
    b = el('tblBorders')
    for side in ('top', 'left', 'bottom', 'right', 'insideH', 'insideV'):
        b.append(el(side, val='single', sz=4, space=0, color='000000'))
    pr.append(b)
    pr.append(el('tblLayout', type='fixed'))
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    for gc, w in zip(tbl.tblGrid.findall(qn('w:gridCol')), tw):
        gc.set(qn('w:w'), str(w))
    for r, row in enumerate(table.rows):
        cells = row.cells
        # dòng tiêu đề nhóm (chỉ ô đầu có chữ) được gộp thành một ô trên cả dòng
        if r > 0 and cells[0].text.strip() and not any(c.text.strip() for c in cells[1:]):
            merged = cells[0].merge(cells[-1])
            for extra in merged.paragraphs[1:]:
                if not extra.text.strip():
                    extra._p.getparent().remove(extra._p)
            tcPr = merged._tc.get_or_add_tcPr()
            for old in tcPr.findall(qn('w:tcW')):
                tcPr.remove(old)
            tcPr.insert(0, el('tcW', w=sum(tw), type='dxa'))
            for p in merged.paragraphs:
                p.paragraph_format.alignment = WD_ALIGN_PARAGRAPH.LEFT
                p.paragraph_format.first_line_indent = Cm(0)
                p.paragraph_format.space_before = Pt(1)
                p.paragraph_format.space_after = Pt(1)
                p.paragraph_format.line_spacing_rule = WD_LINE_SPACING.SINGLE
                for run in p.runs:
                    run.font.size = Pt(11)
                    run.font.name = 'Times New Roman'
            trPr = row._tr.get_or_add_trPr()
            if trPr.find(qn('w:cantSplit')) is None:
                trPr.append(el('cantSplit'))
            continue
        trPr = row._tr.get_or_add_trPr()
        if trPr.find(qn('w:cantSplit')) is None:
            trPr.append(el('cantSplit'))
        if r == 0 and trPr.find(qn('w:tblHeader')) is None:
            trPr.append(el('tblHeader'))
        for c, cell in enumerate(row.cells):
            tcPr = cell._tc.get_or_add_tcPr()
            for old in tcPr.findall(qn('w:tcW')):
                tcPr.remove(old)
            tcPr.insert(0, el('tcW', w=tw[c], type='dxa'))
            for p in cell.paragraphs:
                pf = p.paragraph_format
                pf.first_line_indent = Cm(0)
                pf.left_indent = Cm(0)
                pf.space_before = Pt(1)
                pf.space_after = Pt(1)
                pf.line_spacing_rule = WD_LINE_SPACING.SINGLE
                # dòng cuối đi cùng ghi chú và nguồn của bảng
                pf.keep_with_next = r == len(table.rows) - 1
                if r == 0:
                    pf.alignment = WD_ALIGN_PARAGRAPH.CENTER
                else:
                    pf.alignment = WD_ALIGN_PARAGRAPH.LEFT if c < left_cols else WD_ALIGN_PARAGRAPH.CENTER
                for run in p.runs:
                    run.font.size = Pt(11)
                    run.font.name = 'Times New Roman'
                    run.font.color.rgb = None
                    if r == 0:
                        run.font.bold = True


def page_number_footer(section):
    p = section.footer.paragraphs[0]
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = p.add_run()
    for kind, text in (('begin', None), (None, 'PAGE'), ('end', None)):
        if kind:
            run._r.append(el('fldChar', fldCharType=kind))
        else:
            it = OxmlElement('w:instrText')
            it.set(qn('xml:space'), 'preserve')
            it.text = text
            run._r.append(it)
    run.font.name = 'Times New Roman'
    run.font.size = Pt(11)


def main():
    text = open(SRC, encoding='utf-8').read()
    md, table_ids = to_markdown(text)
    with tempfile.NamedTemporaryFile('w', suffix='.md', delete=False, encoding='utf-8') as f:
        f.write(md)
        tmp = f.name
    try:
        subprocess.run([pandoc_path(), tmp, '-f', 'markdown+pipe_tables-implicit_figures', '-t', 'docx',
                        '--shift-heading-level-by=-1', '-o', OUT], check=True)
    finally:
        os.remove(tmp)
    d = Document(OUT)
    format_styles(d)
    assert len(d.tables) == len(table_ids), (len(d.tables), table_ids)
    for t, num in zip(d.tables, table_ids):
        format_table(t, num)
    for s in d.sections:
        s.page_width, s.page_height = Cm(21), Cm(29.7)
        s.top_margin, s.bottom_margin = Cm(2.5), Cm(2.5)
        s.left_margin, s.right_margin = Cm(3), Cm(2)
        page_number_footer(s)
    # mọi đoạn chữ dùng Times New Roman, kể cả chữ có dấu
    for p in d.paragraphs:
        for run in p.runs:
            run.font.name = 'Times New Roman'
            rpr = run._r.get_or_add_rPr()
            fonts = rpr.find(qn('w:rFonts'))
            if fonts is None:
                fonts = OxmlElement('w:rFonts')
                rpr.insert(0, fonts)
            for a in ('ascii', 'hAnsi', 'cs', 'eastAsia'):
                fonts.set(qn('w:' + a), 'Times New Roman')
    cp = d.core_properties
    cp.author = ''
    cp.last_modified_by = ''
    cp.title = re.search(r'^# (.+)$', text, re.M).group(1).strip()
    d.save(OUT)
    print(OUT)


if __name__ == '__main__':
    main()
