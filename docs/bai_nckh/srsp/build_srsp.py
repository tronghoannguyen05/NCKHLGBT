"""Build the submission files for Sexuality Research and Social Policy (SRSP).

Sources
  docs/bai_nckh/manuscript_en_article.md   masked manuscript, followed by the supplementary material
  docs/bai_nckh/srsp/title_page.md          title page with the identifying information
  docs/bai_nckh/srsp/reference.docx         Word styles (Times New Roman 12, double spacing, 1-inch margins,
                                            page numbers)

Outputs (docs/bai_nckh/srsp/)
  Manuscript.docx              masked: title only on page 1, abstract and keywords, text, Declarations,
                               Compliance with Ethical Standards, Acknowledgments, References, then one
                               table per page, the list of figure captions, and one figure per page
  Title_Page.docx              separate submission item: authors, affiliation, corresponding author,
                               Acknowledgments, Statements and Declarations
  Supplementary_Material.docx  masked supplementary methods and tables

Run: python3 docs/bai_nckh/srsp/build_srsp.py   (requires pandoc or pypandoc_binary, and python-docx)
"""
import copy
import os
import re
import shutil
import subprocess
import tempfile

from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_LINE_SPACING
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt

HERE = os.path.dirname(os.path.abspath(__file__))
BAI = os.path.dirname(HERE)
SRC = os.path.join(BAI, 'manuscript_en_article.md')
TITLE_SRC = os.path.join(HERE, 'title_page.md')
REF = os.path.join(HERE, 'reference.docx')

# Relative column widths, number of left-aligned columns, and page orientation of each table.
TABLES = {
    '1': ([0.80, 0.20], 1, False),
    '2': ([0.70, 0.30], 1, False),
    '3': ([0.28, 0.12, 0.08, 0.08, 0.10, 0.085, 0.085, 0.085, 0.085], 1, True),
    '4': ([0.49, 0.09, 0.15, 0.15, 0.12], 1, False),
    '5': ([0.36, 0.11, 0.11, 0.11, 0.09, 0.11, 0.11], 1, True),
    '6': ([0.38, 0.08, 0.10, 0.10, 0.22, 0.12], 1, False),
    '7': ([0.17, 0.27, 0.12, 0.08, 0.16, 0.10, 0.10], 2, True),
    '8': ([0.50, 0.08, 0.10, 0.20, 0.12], 1, False),
    'S1': ([0.44, 0.14, 0.10, 0.20, 0.12], 1, False),
    'S2': ([0.50, 0.50], 2, False),
    'S3': ([0.55, 0.45], 2, False),
    'S4': ([0.18, 0.36, 0.06, 0.10, 0.14, 0.08, 0.08], 2, True),
    'S5': ([0.17, 0.10, 0.20, 0.14, 0.05, 0.17, 0.17], 4, True),
    'S6': ([0.46, 0.10, 0.10, 0.20, 0.14], 1, False),
}
TEXT_WIDTH = {False: 9360, True: 12960}  # twips: 6.5 in (portrait) and 9 in (landscape)
PAGE_BREAK = '```{=openxml}\n<w:p><w:r><w:br w:type="page"/></w:r></w:p>\n```'


def section_end(landscape):
    """Paragraph that ends the current section; the next section starts on a new page."""
    size = ('<w:pgSz w:w="15840" w:h="12240" w:orient="landscape"/>' if landscape
            else '<w:pgSz w:w="12240" w:h="15840"/>')
    return ('```{=openxml}\n<w:p><w:pPr><w:sectPr><w:type w:val="nextPage"/>' + size +
            '<w:pgMar w:top="1440" w:right="1440" w:bottom="1440" w:left="1440" w:header="720" '
            'w:footer="720" w:gutter="0"/></w:sectPr></w:pPr></w:p>\n```')


def styled(style, text):
    return '::: {custom-style="' + style + '"}\n' + text.strip() + '\n:::'


def pandoc_path():
    try:
        import pypandoc
        return pypandoc.get_pandoc_path()
    except Exception:
        path = shutil.which('pandoc')
        if not path:
            raise SystemExit('pandoc not found')
        return path


def to_docx(markdown, out, shift):
    with tempfile.NamedTemporaryFile('w', suffix='.md', delete=False, encoding='utf-8') as f:
        f.write(markdown)
        tmp = f.name
    try:
        subprocess.run([pandoc_path(), tmp, '-f', 'markdown+pipe_tables-implicit_figures', '-t', 'docx',
                        '--shift-heading-level-by=' + str(shift), '--reference-doc=' + REF,
                        '--resource-path=' + BAI, '-o', out], check=True)
    finally:
        os.remove(tmp)


def split_blocks(text):
    """Separate the running text from the tables and figures, which SRSP wants on separate pages."""
    paras = [p for p in re.split(r'\n\s*\n', text) if p.strip() and p.strip() != '---']
    flow, tables, figures = [], [], []
    i = 0
    while i < len(paras):
        p = paras[i].strip()
        m = re.match(r'\*\*Table (S?\d+)\*\*', p)
        if m:
            assert paras[i + 1].lstrip().startswith('|'), p
            j, notes = i + 2, []
            while j < len(paras) and re.match(r'(\^[a-z]\^|\*Note\.\*)', paras[j].strip()):
                notes.append(paras[j].strip())
                j += 1
            tables.append((m.group(1), p, paras[i + 1].strip(), notes))
            i = j
            continue
        m = re.match(r'\*\*Figure (\d+)\*\* \*(.+)\*$', p)
        if m:
            assert paras[i + 1].lstrip().startswith('!['), p
            note = ''
            if i + 2 < len(paras) and paras[i + 2].strip().startswith('*Note.*'):
                note = paras[i + 2].strip()[len('*Note.*'):].strip()
            figures.append((m.group(1), m.group(2), paras[i + 1].strip(), note))
            i += 3 if note else 2
            continue
        flow.append(paras[i])
        i += 1
    return flow, tables, figures


def table_blocks(tables, first_break):
    """One table per page; wide tables on landscape pages."""
    out, previous_landscape = [], False
    for k, (num, caption, grid, notes) in enumerate(tables):
        landscape = TABLES[num][2]
        if landscape:
            out.append(section_end(previous_landscape))
        elif previous_landscape:
            out.append(section_end(True))
        elif k > 0 or first_break:
            out.append(PAGE_BREAK)
        out.append(styled('Table Title', caption))
        out.append(grid)
        out.extend(styled('Table Note', n) for n in notes)
        previous_landscape = landscape
    if previous_landscape:
        out.append(section_end(True))
    return out, previous_landscape


def set_child(parent, child, order):
    """Insert child into parent at the position required by the schema, replacing an element of the same tag."""
    tag = child.tag.split('}')[1]
    for old in parent.findall(qn('w:' + tag)):
        parent.remove(old)
    later = order[order.index(tag) + 1:]
    anchor = next((c for c in parent if c.tag.split('}')[1] in later), None)
    if anchor is not None:
        anchor.addprevious(child)
    else:
        parent.append(child)


TBLPR = ['tblStyle', 'tblpPr', 'tblOverlap', 'bidiVisual', 'tblStyleRowBandSize', 'tblStyleColBandSize', 'tblW',
         'jc', 'tblCellSpacing', 'tblInd', 'tblBorders', 'shd', 'tblLayout', 'tblCellMar', 'tblLook',
         'tblCaption', 'tblDescription']
TCPR = ['cnfStyle', 'tcW', 'gridSpan', 'hMerge', 'vMerge', 'tcBorders', 'shd', 'noWrap', 'tcMar',
        'textDirection', 'tcFitText', 'vAlign', 'hideMark']
PPR = ['pStyle', 'keepNext', 'keepLines', 'pageBreakBefore', 'framePr', 'widowControl', 'numPr',
       'suppressLineNumbers', 'pBdr', 'shd', 'tabs', 'suppressAutoHyphens', 'kinsoku', 'wordWrap',
       'overflowPunct', 'topLinePunct', 'autoSpaceDE', 'autoSpaceDN', 'bidi', 'adjustRightInd', 'snapToGrid',
       'spacing', 'ind', 'contextualSpacing', 'mirrorIndents', 'suppressOverlap', 'jc', 'textDirection',
       'textAlignment', 'textboxTightWrap', 'outlineLvl', 'divId', 'cnfStyle', 'rPr', 'sectPr', 'pPrChange']
TRPR = ['cnfStyle', 'divId', 'gridBefore', 'gridAfter', 'wBefore', 'wAfter', 'cantSplit', 'trHeight',
        'tblHeader', 'tblCellSpacing', 'jc', 'hidden']


def el(tag, **attrs):
    e = OxmlElement('w:' + tag)
    for k, v in attrs.items():
        e.set(qn('w:' + k), str(v))
    return e


def borders(tag, sides):
    b = el(tag)
    for side in ('top', 'left', 'bottom', 'right', 'insideH', 'insideV'):
        if side in sides:
            b.append(el(side, val=sides[side], sz=6 if sides[side] != 'nil' else 0, space=0, color='000000'))
    return b


def format_table(table, num):
    widths, left_cols, landscape = TABLES[num]
    ncol = len(table.columns)
    assert ncol == len(widths), (num, ncol, len(widths))
    total = TEXT_WIDTH[landscape]
    tw = [int(round(w * total)) for w in widths]
    tbl = table._tbl
    pr = tbl.tblPr
    set_child(pr, el('tblW', w=sum(tw), type='dxa'), TBLPR)
    set_child(pr, el('jc', val='center'), TBLPR)
    set_child(pr, borders('tblBorders', {'top': 'single', 'bottom': 'single', 'left': 'nil', 'right': 'nil',
                                         'insideH': 'nil', 'insideV': 'nil'}), TBLPR)
    set_child(pr, el('tblLayout', type='fixed'), TBLPR)
    for gc, w in zip(tbl.tblGrid.findall(qn('w:gridCol')), tw):
        gc.set(qn('w:w'), str(w))
    for r, row in enumerate(table.rows):
        trPr = row._tr.get_or_add_trPr()
        set_child(trPr, el('cantSplit'), TRPR)
        if r == 0:
            set_child(trPr, el('tblHeader'), TRPR)
        for c, cell in enumerate(row.cells):
            tcPr = cell._tc.get_or_add_tcPr()
            set_child(tcPr, el('tcW', w=tw[c], type='dxa'), TCPR)
            if r == 0:
                set_child(tcPr, borders('tcBorders', {'bottom': 'single'}), TCPR)
            for p in cell.paragraphs:
                pf = p.paragraph_format
                pf.first_line_indent = Inches(0)
                pf.left_indent = Inches(0)
                pf.space_before = Pt(0)
                pf.space_after = Pt(0)
                pf.line_spacing_rule = WD_LINE_SPACING.DOUBLE
                pf.alignment = WD_ALIGN_PARAGRAPH.LEFT if c < left_cols else WD_ALIGN_PARAGRAPH.CENTER
                for run in p.runs:
                    run.font.size = Pt(10)
                    run.font.name = 'Times New Roman'


def style_paragraphs(d):
    def para_style(name, align=None, indent=0.0, hanging=None, size=None, bold=None, keep=False):
        try:
            st = d.styles[name]
        except KeyError:
            return
        pf = st.paragraph_format
        pf.space_before = Pt(0)
        pf.space_after = Pt(0)
        pf.line_spacing_rule = WD_LINE_SPACING.DOUBLE
        if align is not None:
            pf.alignment = align
        if hanging is not None:
            pf.left_indent = Inches(hanging)
            pf.first_line_indent = Inches(-hanging)
        else:
            pf.first_line_indent = Inches(indent)
        pf.keep_with_next = keep
        st.font.name = 'Times New Roman'
        if size:
            st.font.size = Pt(size)
        if bold is not None:
            st.font.bold = bold

    para_style('Title', align=WD_ALIGN_PARAGRAPH.CENTER, size=14, bold=True)
    para_style('Subtitle Line', align=WD_ALIGN_PARAGRAPH.CENTER, size=12)
    para_style('Table Title', align=WD_ALIGN_PARAGRAPH.CENTER, size=12, keep=True)
    para_style('Table Note', align=WD_ALIGN_PARAGRAPH.LEFT, size=10)
    para_style('Abstract Text', align=WD_ALIGN_PARAGRAPH.LEFT, size=12)
    para_style('Reference Entry', align=WD_ALIGN_PARAGRAPH.LEFT, hanging=0.5, size=12)
    para_style('Figure Caption', align=WD_ALIGN_PARAGRAPH.LEFT, size=12)
    para_style('Figure Label', align=WD_ALIGN_PARAGRAPH.LEFT, size=12, bold=True, keep=True)
    para_style('Figure Image', align=WD_ALIGN_PARAGRAPH.CENTER)
    para_style('Plain Block', align=WD_ALIGN_PARAGRAPH.LEFT, size=12)


def page_breaks_to_property(d):
    """Replace each empty page-break paragraph by 'page break before' on the next paragraph.

    A separate break paragraph that falls at the top of a page would otherwise leave a blank page."""
    body = d.element.body
    for p in list(body.iterchildren(qn('w:p'))):
        brs = p.findall('.//' + qn('w:br'))
        text = ''.join(t.text or '' for t in p.iter(qn('w:t')))
        if (len(brs) != 1 or brs[0].get(qn('w:type')) != 'page' or text.strip()
                or p.find('.//' + qn('w:sectPr')) is not None):
            continue
        nxt = p.getnext()
        if nxt is None or nxt.tag != qn('w:p'):
            continue
        pPr = nxt.find(qn('w:pPr'))
        if pPr is None:
            pPr = OxmlElement('w:pPr')
            nxt.insert(0, pPr)
        set_child(pPr, el('pageBreakBefore'), PPR)
        body.remove(p)


def finish(path, title, table_ids):
    d = Document(path)
    style_paragraphs(d)
    page_breaks_to_property(d)
    assert len(d.tables) == len(table_ids), (path, len(d.tables), table_ids)
    for t, num in zip(d.tables, table_ids):
        format_table(t, num)
    # every section keeps the page-number footer and the same margins
    body_sect = d.sections[-1]._sectPr
    refs = body_sect.findall(qn('w:footerReference')) + body_sect.findall(qn('w:headerReference'))
    for sp in d.element.body.iter(qn('w:sectPr')):
        if sp is body_sect:
            continue
        if not sp.findall(qn('w:footerReference')):
            for ref in reversed(refs):
                sp.insert(0, copy.deepcopy(ref))
    sect = d.sections[-1]
    for s in d.sections:
        s.left_margin = s.right_margin = s.top_margin = s.bottom_margin = Inches(1)
    sect.page_width, sect.page_height = Inches(8.5), Inches(11)
    cp = d.core_properties
    cp.author = ''
    cp.last_modified_by = ''
    cp.comments = ''
    cp.keywords = ''
    cp.subject = ''
    cp.category = ''
    cp.title = title
    d.save(path)


def manuscript(text, title):
    main = text.split('\n## Supplementary Material\n')[0]
    body, refs = main.split('\n## References\n')
    body = re.sub(r'^# .+\n', '', body, count=1)
    flow, tables, figures = split_blocks(body)
    out = [styled('Title', title), PAGE_BREAK]
    in_abstract = False
    for p in flow:
        s = p.strip()
        if s.startswith('## Abstract'):
            in_abstract = True
            out.append(s)
            continue
        if s.startswith('## Introduction'):
            in_abstract = False
            out.append(PAGE_BREAK)
        out.append(styled('Abstract Text', s) if in_abstract else s)
    out.append(PAGE_BREAK)
    out.append('## References')
    out.extend(styled('Reference Entry', r) for r in re.split(r'\n\s*\n', refs) if r.strip() and r.strip() != '---')
    blocks, ended_section = table_blocks(tables, first_break=True)
    out.extend(blocks)
    if not ended_section:
        out.append(PAGE_BREAK)
    out.append('## Figure Captions')
    for num, cap, img, note in figures:
        out.append(styled('Figure Caption', '**Figure ' + num + '** *' + cap.rstrip('.') + '.* ' + note))
    for num, cap, img, note in figures:
        out.append(PAGE_BREAK)
        out.append(styled('Figure Label', 'Figure ' + num))
        out.append(styled('Figure Image', re.sub(r'\)$', '){width=6.0in}', img)))
    return '\n\n'.join(out) + '\n', [t[0] for t in tables]


def supplementary(text, title):
    supp = text.split('\n## Supplementary Material\n')[1]
    flow_part, tables_part = supp.split('\n### Supplementary Tables\n')
    out = [styled('Title', 'Supplementary Material'), styled('Subtitle Line', '*' + title + '*')]
    for p in (p.strip() for p in re.split(r'\n\s*\n', flow_part) if p.strip()):
        out.append(p if p.startswith(('#', '**S')) else styled('Subtitle Line', p))
    out.append(PAGE_BREAK)
    out.append('### Supplementary Tables')
    flow, tables, _ = split_blocks(tables_part)
    blocks, _ = table_blocks(tables, first_break=False)
    out.extend(blocks)
    return '\n\n'.join(out) + '\n', [t[0] for t in tables]


def title_page(title):
    text = open(TITLE_SRC, encoding='utf-8').read()
    text = re.sub(r'^# .+\n', '', text, count=1)
    paras = [p.strip() for p in re.split(r'\n\s*\n', text) if p.strip()]
    out = [styled('Title', title), styled('Subtitle Line', paras[0])]
    for p in paras[1:]:
        if p.startswith('#') or p.startswith('[^'):
            out.append(p)
        else:
            out.append(styled('Plain Block', p))
    return '\n\n'.join(out) + '\n'


def main():
    text = open(SRC, encoding='utf-8').read()
    title = re.search(r'^# (.+)$', text, re.M).group(1).strip()

    md, ids = manuscript(text, title)
    out = os.path.join(HERE, 'Manuscript.docx')
    to_docx(md, out, -1)
    finish(out, title, ids)
    print(out)

    md, ids = supplementary(text, title)
    out = os.path.join(HERE, 'Supplementary_Material.docx')
    to_docx(md, out, -2)
    finish(out, 'Supplementary Material', ids)
    print(out)

    out = os.path.join(HERE, 'Title_Page.docx')
    to_docx(title_page(title), out, -1)
    finish(out, title, [])
    print(out)


if __name__ == '__main__':
    main()
